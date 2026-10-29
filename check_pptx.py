import os
os.chdir("D:\python_p\Adaptive_Learning_Project")
from pptx import Presentation
prs = Presentation("Project_Presentation_Adaptive_Learning.pptx")
stale_keywords = ['56.07', '69.58', '54.94', '62.92', 'Logistic', 'Random Forest', 'Accuracy', 'macro']
print("=== PPTX CHECK ===")
for i, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            t = shape.text_frame.text.strip()
            for kw in stale_keywords:
                if kw in t:
                    print("[slide {}] contains '{}': {}".format(i, kw, t[:100]))
                    break
        if shape.has_table:
            for ri, row in enumerate(shape.table.rows):
                for ci, cell in enumerate(row.cells):
                    t = cell.text.strip()
                    for kw in stale_keywords:
                        if kw in t:
                            print("[slide {} table row {} cell {}] contains '{}': {}".format(i, ri, ci, kw, t[:100]))
                            break
print("Done.")