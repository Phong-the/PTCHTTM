import sys
import os
os.chdir("D:\python_p\Adaptive_Learning_Project")

from docx import Document

doc = Document("Project_Action_Plan_Adaptive_Learning.docx")
print("=== ACTION PLAN DOCX ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t and any(kw in t for kw in ['56.07', '69.58', '54.94', '62.92', 'Logistic', 'Random Forest', 'Accuracy', 'macro']):
        print("[P{}] {}".format(i, t.encode('ascii', 'replace').decode('ascii')))

for ti, table in enumerate(doc.tables):
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        joined = " || ".join(cells)
        if any(kw in joined for kw in ['56.07', '69.58', '54.94', '62.92', 'Logistic', 'Random Forest', 'Accuracy', 'macro']):
            print("[T{}R{}] {}".format(ti, ri, joined.encode('ascii', 'replace').decode('ascii')))

print("\n=== FINAL REPORT DOCX ===")
doc2 = Document("Project_Final_Report_Adaptive_Learning.docx")
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    if t and any(kw in t for kw in ['56.07', '69.58', '54.94', '62.92', '0.5607', 'Logistic', 'Random Forest', 'Accuracy', 'macro', 'MSSV', 'Nhom pho', 'Thuong nhom', '__________', 'Projected Impact', 'Ket luc', 'Future Improvements', 'Accomplishments']):
        print("[P{}] {}".format(i, t.encode('ascii', 'replace').decode('ascii')))

for ti, table in enumerate(doc2.tables):
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        joined = " || ".join(cells)
        if any(kw in joined for kw in ['56.07', '69.58', '54.94', '62.92', '0.5607', 'Logistic', 'Random Forest', 'Accuracy', 'macro']):
            print("[T{}R{}] {}".format(ti, ri, joined.encode('ascii', 'replace').decode('ascii')))

print("\n=== PPTX ===")
from pptx import Presentation
prs = Presentation("Project_Presentation_Adaptive_Learning.pptx")
for si, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            t = shape.text_frame.text.strip()
            t_ascii = t.encode('ascii', 'replace').decode('ascii')
            if any(kw in t_ascii for kw in ['56.07', '69.58', '54.94', '62.92', 'Logistic', 'Random Forest', 'Accuracy', 'macro']):
                print("[slide {}] {}".format(si, t_ascii))